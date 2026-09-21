#!/usr/bin/env python3
"""
Import monthly data from Expences.xlsx into the personal finance app.

Usage:
    python import-excel.py --sheet stanislav --token <JWT> [--file <path>]
    python import-excel.py --sheet lea       --token <JWT> [--file <path>]
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl
import requests

BASE = "http://localhost:8080/api"
CURRENCY = "CZK"
CHUNK = 25

MONTH_MAP = {
    "Január": 1, "Február": 2, "Marec": 3, "Apríl": 4,
    "Máj": 5, "Jún": 6, "Júl": 7, "August": 8,
    "September": 9, "Október": 10, "November": 11, "December": 12,
    # typos present in the sheet
    "Nobember": 11, "Octóber": 10, "April": 4,
}

# column name -> (category name, type, necessity)
STANISLAV_COLS = {
    "Salary":       ("Salary",       "INCOME",     "NEED"),
    "Gastro":       ("Salary",       "INCOME",     "NEED"),
    "Flexi":        ("Salary",       "INCOME",     "NEED"),
    "Rent":         ("Rent",         "EXPENSE",    "NEED"),
    "Energie":      ("Energy",       "EXPENSE",    "NEED"),
    "Elektricity":  ("Electricity",  "EXPENSE",    "NEED"),
    "Internet":     ("Internet",     "EXPENSE",    "NEED"),
    "Phone":        ("Phone",        "EXPENSE",    "NEED"),
    "Insurance":    ("Insurance",    "EXPENSE",    "NEED"),
    "Groceries":    ("Groceries",    "EXPENSE",    "NEED"),
    "Household":    ("Household",    "EXPENSE",    "NEED"),
    "Transport":    ("Transport",    "EXPENSE",    "NEED"),
    "Clothing":     ("Clothing",     "EXPENSE",    "WANT"),
    "Multisport":   ("Subscription", "EXPENSE",    "WANT"),
    "Subscription": ("Subscription", "EXPENSE",    "WANT"),
    "Restaurants":  ("Restaurants",  "EXPENSE",    "WANT"),
    "Alza":         ("Alza",         "EXPENSE",    "WANT"),
    "Entertainment":("Entertainment","EXPENSE",    "WANT"),
    "Other":        ("Other",        "EXPENSE",    "WANT"),
    "ETF":          ("Stocks",       "INVESTMENT", "NEED"),
    "Indepedence":  ("Savings",      "INVESTMENT", "NEED"),
}

LEA_COLS = {
    "Salary":       ("Salary",       "INCOME",     "NEED"),
    "Gastro":       ("Salary",       "INCOME",     "NEED"),
    "Rent":         ("Rent",         "EXPENSE",    "NEED"),
    "Phone":        ("Phone",        "EXPENSE",    "NEED"),
    "Groceries":    ("Groceries",    "EXPENSE",    "NEED"),
    "Household":    ("Household",    "EXPENSE",    "NEED"),
    "Transport":    ("Transport",    "EXPENSE",    "NEED"),
    "Clothing":     ("Clothing",     "EXPENSE",    "WANT"),
    "Restaurants":  ("Restaurants",  "EXPENSE",    "WANT"),
    "Alza":         ("Alza",         "EXPENSE",    "WANT"),
    "Entertainment":("Entertainment","EXPENSE",    "WANT"),
    "Other":        ("Other",        "EXPENSE",    "WANT"),
    "ETF":          ("Stocks",       "INVESTMENT", "NEED"),
    "Car":          ("Car",          "EXPENSE",    "NEED"),
    "Hypotéka":     ("Mortgage",     "EXPENSE",    "NEED"),
}


def fetch_category_map(token: str) -> dict:
    """Returns {category_name: categoryId}."""
    resp = requests.get(
        f"{BASE}/categories",
        headers={"Authorization": f"Bearer {token}"},
    )
    resp.raise_for_status()
    categories = resp.json()
    return {c["name"]: c["categoryId"] for c in categories}


def parse_value(raw):
    """Return float or None; skip zero and non-numeric."""
    if raw is None:
        return None
    try:
        v = float(raw)
    except (TypeError, ValueError):
        return None
    return v if v != 0 else None


def build_entries(ws, col_map: dict, cat_id_map: dict) -> list:
    headers = [cell.value for cell in ws[1]]
    entries = []
    skipped_cats = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        month_name = row[0]
        year = row[1]

        # Skip totals row and empty rows
        if not month_name or month_name == "Total":
            continue
        if month_name not in MONTH_MAP:
            print(f"  WARNING: unknown month '{month_name}' — skipping row", file=sys.stderr)
            continue
        if not year:
            continue

        month_num = MONTH_MAP[month_name]
        date_str = f"{int(year)}-{month_num:02d}-01"

        for col_idx, header in enumerate(headers):
            if header not in col_map:
                continue  # metadata or sum column

            cat_name, entry_type, necessity = col_map[header]
            value = parse_value(row[col_idx])
            if value is None:
                continue  # zero or null — skip

            cat_id = cat_id_map.get(cat_name)
            if cat_id is None:
                if cat_name not in skipped_cats:
                    print(f"  WARNING: category '{cat_name}' not found in app — entries skipped", file=sys.stderr)
                    skipped_cats.add(cat_name)
                continue

            entries.append({
                "amount": {"value": value, "currency": CURRENCY},
                "categoryId": cat_id,
                "date": date_str,
                "name": header,
                "note": "",
                "type": entry_type,
                "necessity": necessity,
            })

    return entries


def send_entries(entries: list, token: str) -> None:
    total = len(entries)
    print(f"  Sending {total} entries in chunks of {CHUNK}...")
    for i in range(0, total, CHUNK):
        chunk = entries[i : i + CHUNK]
        resp = requests.post(
            f"{BASE}/entries/batch",
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/json",
            },
            data=json.dumps(chunk),
        )
        batch_num = i // CHUNK + 1
        if resp.status_code == 200:
            print(f"  Batch {batch_num}: OK ({len(chunk)} entries)")
        else:
            print(f"  Batch {batch_num}: ERROR {resp.status_code} — {resp.text}", file=sys.stderr)


def main():
    parser = argparse.ArgumentParser(description="Import Expences.xlsx into the finance app")
    parser.add_argument("--sheet", required=True, choices=["stanislav", "lea"])
    parser.add_argument("--token", required=True, help="JWT access token (without 'Bearer ' prefix)")
    parser.add_argument("--file", default="Expences.xlsx", help="Path to the Excel file")
    args = parser.parse_args()

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        print(f"ERROR: file not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    sheet_name = args.sheet.capitalize()  # "Stanislav" or "Lea"
    col_map = STANISLAV_COLS if args.sheet == "stanislav" else LEA_COLS

    print(f"Loading {xlsx_path} — sheet '{sheet_name}'...")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: sheet '{sheet_name}' not found. Available: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)
    ws = wb[sheet_name]

    print("Fetching categories from API...")
    cat_id_map = fetch_category_map(args.token)
    print(f"  Found {len(cat_id_map)} categories: {list(cat_id_map.keys())}")

    print("Building entries...")
    entries = build_entries(ws, col_map, cat_id_map)
    print(f"  Built {len(entries)} entries (zeros skipped)")

    send_entries(entries, args.token)
    print("Done.")


if __name__ == "__main__":
    main()
